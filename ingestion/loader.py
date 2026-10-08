import pdfplumber
from pathlib import Path
from docx import Document

#==================================== Document Loader ==========================================
# checks file suffix for it's type
def load_document(file_path : str) -> str :
    path = Path(file_path)

    if not path.exists() :
        raise FileNotFoundError(f"file not found : {file_path}")

    if path.suffix.lower() == ".pdf" :
        return load_pdf(file_path)
    elif path.suffix.lower() in [".txt", ".md"] :
        return path.read_text(encoding="utf-8")
    elif path.suffix.lower() == ".docx" :
        return load_docx(file_path)
    else :
        raise ValueError(f"Unsupported file type : {path.suffix}")

#==================================== PDF Loader ==========================================
def load_pdf(file_path : str) -> str :
    all_content = []

    with pdfplumber.open(file_path) as pdf :
        for page in pdf.pages :

            # Extract Plain text from page
            page_text = page.extract_text()
            if page_text :
                all_content.append(page_text)

            # Extract tables form page
            tables = page.extract_tables()
            for table in tables :
                table_as_text = convert_table_to_text(table)
                all_content.append(table_as_text)

    # Merge Everything
    full_text = "\n".join(all_content)
    return full_text.strip()

#==================================== Docx Loader ==========================================
def load_docx(file_path : str) -> str :
    doc = Document(file_path)
    all_content = []

    # Extract paragraphs 
    for paragraph in doc.paragraphs :
        if paragraph.text.strip() :
            all_content.append(paragraph.text.strip())

    # Extract tables
    for table in doc.tables :
        rows = []
        for row in table.rows :
            cell_values = [cell.text.strip() for cell in row.cells]
            rows.append(cell_values)

        table_as_text = convert_table_to_text(rows)
        all_content.append(table_as_text)

    full_text = "\n".join(all_content)
    return full_text.strip()

#==================================== Row Extractor ==========================================
def convert_table_to_text(table : list) -> str :
    all_rows = []
    for row in table :
        cleaned_cells = []
        for cell in row :
            cell_text = str(cell).strip() if cell else ""
            cleaned_cells.append(cell_text)

        row_as_text = " | ".join(cleaned_cells)
        all_rows.append(row_as_text)

    return "\n".join(all_rows)

#=============================================================================================